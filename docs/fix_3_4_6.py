# -*- coding: utf-8 -*-
"""Rewrite section 3.4.6 - fix orphaned tables and rebuild content."""
import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph
from docx.table import Table

BASE = os.path.dirname(os.path.abspath(__file__))
DOC = os.path.join(BASE, '\u7cfb\u7edf\u8bbe\u8ba1\u4e0e\u7b97\u6cd5\u62a5\u544a .docx')
CHARTS = os.path.join(BASE, 'charts')

def make_table(doc, headers, rows):
    """Create a well-formed table element with proper XML structure."""
    tbl = OxmlElement('w:tbl')

    # tblPr
    tblPr = OxmlElement('w:tblPr')
    tblW = OxmlElement('w:tblW')
    tblW.set(qn('w:w'), '5000')
    tblW.set(qn('w:type'), 'pct')
    tblPr.append(tblW)
    jc = OxmlElement('w:jc')
    jc.set(qn('w:val'), 'center')
    tblPr.append(jc)
    tblBorders = OxmlElement('w:tblBorders')
    for name in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        b = OxmlElement('w:' + name)
        b.set(qn('w:val'), 'single')
        b.set(qn('w:sz'), '4')
        b.set(qn('w:space'), '0')
        b.set(qn('w:color'), 'auto')
        tblBorders.append(b)
    tblPr.append(tblBorders)
    tbl.append(tblPr)

    # tblGrid
    tblGrid = OxmlElement('w:tblGrid')
    for _ in headers:
        gc = OxmlElement('w:gridCol')
        tblGrid.append(gc)
    tbl.append(tblGrid)

    # Wrap and add rows
    table = Table(tbl, doc)

    # Header row
    row = table.add_row()
    for i, h in enumerate(headers):
        c = row.cells[i]
        c.text = h
        for p in c.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                r.bold = True
                r.font.size = Pt(9)

    # Data rows
    for row_data in rows:
        row = table.add_row()
        for i, val in enumerate(row_data):
            c = row.cells[i]
            c.text = str(val)
            for p in c.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for r in p.runs:
                    r.font.size = Pt(9)

    return table, tbl

def insert_after(ref_elem, new_elem):
    """Insert new_elem after ref_elem in the parent."""
    ref_elem.addnext(new_elem)

def add_para(doc, ref_elem, text='', bold_parts=None):
    p_elem = OxmlElement('w:p')
    insert_after(ref_elem, p_elem)
    p = Paragraph(p_elem, doc)
    try:
        p.style = doc.styles['Normal']
    except:
        pass
    if bold_parts:
        for part, is_bold in bold_parts:
            r = p.add_run(part)
            r.bold = is_bold
            r.font.size = Pt(10.5)
    elif text:
        r = p.add_run(text)
        r.font.size = Pt(10.5)
    return p, p_elem

def add_heading(doc, ref_elem, text, level=3):
    p_elem = OxmlElement('w:p')
    insert_after(ref_elem, p_elem)
    p = Paragraph(p_elem, doc)
    try:
        p.style = doc.styles['Heading %d' % level]
    except:
        pass
    p.add_run(text)
    return p, p_elem

def add_bullet(doc, ref_elem, text):
    p_elem = OxmlElement('w:p')
    insert_after(ref_elem, p_elem)
    p = Paragraph(p_elem, doc)
    try:
        p.style = doc.styles['List Bullet']
    except:
        r = p.add_run('\u2022 ' + text)
        r.font.size = Pt(10.5)
        return p, p_elem
    r = p.add_run(text)
    r.font.size = Pt(10.5)
    return p, p_elem

def add_para_with_table(doc, ref_elem, text_before, headers, rows, caption=None):
    """Add a paragraph, then a table, then an optional caption paragraph.
    Each element is properly separated as a distinct XML sibling."""
    # 1. Text paragraph
    p1, e1 = add_para(doc, ref_elem, text_before)

    # 2. Table (as a sibling, not nested)
    table, tbl = make_table(doc, headers, rows)
    insert_after(e1, tbl)

    # 3. Caption paragraph after table (acts as separator)
    if caption:
        cp, ce = add_para(doc, tbl, caption)
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in cp.runs:
            r.font.size = Pt(9)
            r.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
        return cp, ce
    else:
        # Add an empty paragraph as separator after table
        sp, se = add_para(doc, tbl, '')
        return sp, se

def main():
    print('Opening document...')
    doc = Document(DOC)
    body = doc.element.body
    elements = list(body)

    # Find 3.4.6 heading
    target_idx = None
    ch4_idx = None
    for i, elem in enumerate(elements):
        if elem.tag == qn('w:p'):
            text = ''.join(t.text or '' for t in elem.iter(qn('w:t')))
            if '3.4.6' in text and '\u76ee\u6807\u8fbe\u6210' in text:
                target_idx = i
                print(f'Found 3.4.6 at element index {i}')
            elif target_idx is not None and text.strip().startswith('4.') and '\u603b\u7ed3' in text:
                ch4_idx = i
                print(f'Found Chapter 4 at element index {i}')
                break

    if target_idx is None or ch4_idx is None:
        print('ERROR: Could not find section boundaries')
        return

    # Delete ALL elements (paragraphs AND tables) between 3.4.6 and Chapter 4
    print(f'Deleting elements {target_idx + 1} to {ch4_idx - 1}...')
    to_delete = elements[target_idx + 1 : ch4_idx]
    deleted_count = 0
    for elem in to_delete:
        body.remove(elem)
        deleted_count += 1
    print(f'Deleted {deleted_count} elements (including orphaned tables)')

    # Re-read elements after deletion
    elements = list(body)
    target_elem = elements[target_idx]

    # Build new 3.4.6 content
    ref = target_elem

    # --- Section text: goal achievement summary ---
    ref = add_para(doc, ref, '', bold_parts=[
        ('\u5de5\u7a0b\u76ee\u6807\u9010\u9879\u6838\u5bf9\u5982\u4e0b\u8868\u6240\u793a\uff1a', False),
    ])[1]

    # Table 3-14: Engineering goals checklist
    ref = add_para_with_table(doc, ref,
        '',
        ['\u5e8f\u53f7', '\u5de5\u7a0b\u76ee\u6807', '\u8d5b\u9898\u8981\u6c42', '\u5b9e\u6d4b\u7ed3\u679c', '\u5224\u5b9a'],
        [
            ['1', '\u8def\u7f51\u89c4\u6a21', '\u226520\u8def\u53e3', '30\u8def\u53e36\u00d75\u7f51\u683c', '\u2705 \u8d85\u989d\u5b8c\u6210'],
            ['2', '\u72b6\u6001\u5951\u7ea6', '\u6807\u51c6\u5316\u89c2\u6d4b', '660\u7ef4\u5168\u5c40\u72b6\u6001 + 22\u7ef4\u5c40\u90e8\u89c2\u6d4b', '\u2705 \u8fbe\u6807'],
            ['3', '\u51b3\u7b56\u7ea6\u675f', '\u5b89\u5168\u52a8\u4f5c\u7a7a\u95f4', '4\u52a8\u4f5c + \u9700\u6c42\u95e8\u63a7\u63a9\u7801\u5951\u7ea6', '\u2705 \u8fbe\u6807'],
            ['4', '\u6a21\u578b\u4f53\u79ef', '< 1 MB', '25.14 KB (FP32 ONNX)', '\u2705 \u8fdc\u8d85\u76ee\u6807'],
            ['5', '\u5355\u8def\u53e3\u63a8\u7406\u5ef6\u8fdf', '< 5 ms', '0.022 ms (mean) / 0.028 ms (P95)', '\u2705 \u8fdc\u8d85\u76ee\u6807'],
            ['6', '30\u8def\u53e3\u6279\u91cf\u5ef6\u8fdf', '< 500 ms', '0.029 ms (mean) / 0.048 ms (P95)', '\u2705 \u8fdc\u8d85\u76ee\u6807'],
            ['7', '\u52a8\u4f5c\u9000\u5316\u7387', '\u22645%', '0%\uff08\u771f\u5b9e\u72b6\u6001100%\u4e00\u81f4\u7387\uff09', '\u2705 \u8fbe\u6807'],
            ['8', '\u4e09\u573a\u666f\u5bf9\u6bd4\u8bc4\u4f30', '\u540c\u8d1f\u8377\u5bf9\u6bd4', '\u65e9/\u665a\u9ad8\u5cf0reward +11~15%\uff0c30\u8def\u53e320/30\u80dc', '\u2705 \u4e3b\u4f53\u8fbe\u6807'],
            ['9', '\u6f14\u793a\u95ed\u73af', '\u4eff\u771f+\u63a8\u7406+\u53ef\u89c6\u5316', 'Unity\u2013WebSocket\u2013SUMO\u2013ONNX\u4e09\u573a\u666f\u95ed\u73af', '\u2705 \u8fbe\u6807'],
            ['10', 'LLM\u4e91\u8111\uff08\u8d5b\u9053C\uff09', '\u4e8b\u4ef6\u8bc6\u522b+\u5efa\u8bae', 'f32 GGUF\u51c6\u786e\u7387100%\uff0c\u5ef6\u8fdf~5s', '\u2705 \u8fbe\u6807'],
        ],
        '\u88683-14 \u5de5\u7a0b\u76ee\u6807\u8d5b\u9898\u8fbe\u6807\u6838\u5bf9'
    )[1]

    # --- Negative findings summary ---
    ref = add_para(doc, ref, '', bold_parts=[
        ('\u6709\u4ef7\u503c\u7684\u8d1f\u9762\u53d1\u73b0\uff1a', True),
        ('\u91cf\u5316\u624b\u6bb5\u5fc5\u987b\u9010\u4efb\u52a1\u5b9e\u6d4b\u9a8c\u8bc1\u3002INT8\u52a8\u6001\u91cf\u5316\u5728\u771f\u5b9eSUMO\u8f68\u8ff9\u4e0a\u4e00\u81f4\u7387\u4ec584\u201393%\uff08\u79bb\u7ebf99.3%\uff09\uff0c\u526a\u679d\u4ea7\u7269\u79bb\u7ebf\u4e00\u81f4\u7387\u7ea696.7%\uff0c\u5747\u672a\u8fbe\u5230\u4e0eFP32 ONNX\u76f8\u540c\u7684\u90e8\u7f72\u95e8\u69db\u3002LLM\u4e91\u8111\u7684INT4/INT6\u91cf\u5316\u5bfc\u81f4\u4e8b\u4ef6\u53ec\u56de\u4ece100%\u96f6\u52300\u20134%\uff0c0.5B\u5c0f\u6a21\u578b\u5fc5\u987b\u4f7f\u7528f32\u8ba1\u7b97\u3002\u8fd9\u4e9b\u201c\u7406\u60f3\u72b6\u6001\u8868\u73b0\u597d\u3001\u771f\u5b9e\u90e8\u7f72\u5931\u6548\u201d\u7684\u73b0\u8c61\uff0c\u6b63\u662f\u8d5b\u9898\u70b9\u540d\u7684\u884c\u4e1a\u75db\u70b9\u3002', False),
    ])[1]

    # --- Table 3-15: Lightweighting comparison ---
    ref = add_para_with_table(doc, ref,
        '',
        ['\u8f7b\u91cf\u5316\u65b9\u6848', '\u6a21\u578b\u4f53\u79ef', '\u79bb\u7ebf\u4e00\u81f4\u7387', '\u771f\u5b9e\u72b6\u6001\u4e00\u81f4\u7387', '\u5224\u5b9a'],
        [
            ['FP32 ONNX\uff08\u6b63\u5f0f\u90e8\u7f72\uff09', '25.14 KB', '100%', '100%', '\u2705 \u90e8\u7f72\u8fbe\u6807'],
            ['INT8 \u52a8\u6001\u91cf\u5316', '~8 KB', '99.3%', '84\u201393%', '\u274c \u771f\u5b9e\u5931\u6548'],
            ['\u7ed3\u6784\u5316\u526a\u679d 64\u219256', '~20 KB', '96.7%', '\u672a\u5b9e\u6d4b', '\u26a0\ufe0f \u5019\u9009/\u6d88\u878d'],
        ],
        '\u88683-15 DQN\u8f7b\u91cf\u5316\u65b9\u6848\u5bf9\u6bd4'
    )[1]

    # --- Table 3-16: LLM quantization comparison ---
    ref = add_para_with_table(doc, ref,
        '',
        ['\u6a21\u578b\u5f62\u6001', '\u4f53\u79ef', '\u4e8b\u4ef6\u51c6\u786e\u7387', '\u4e8b\u4ef6\u53ec\u56de', '\u5355\u6761\u5ef6\u8fdf', '\u5224\u5b9a'],
        [
            ['\u57fa\u5ea7bf16\uff08\u5fae\u8c03\u524d\uff09', '~1 GB', '25.0%', '0%', '-', '\u57fa\u7ebf'],
            ['\u5fae\u8c03\u540ebf16\uff08HF\u63a8\u7406\uff09', '~1 GB', '99.5%', '100%', '~6.1 s', '\u2705'],
            ['\u5fae\u8c03\u540ef32 GGUF\uff08\u6b63\u5f0f\uff09', '1.98 GB', '100.0%', '100%', '~5.0 s (P95 6.2s)', '\u2705 \u6b63\u5f0f\u90e8\u7f72'],
            ['Q6_K', '506 MB', '66.0%', '4%', '~3.2 s', '\u274c \u91cf\u5316\u5931\u6548'],
            ['Q4_K_M', '379 MB', '62.0%', '0%', '~3.2 s', '\u274c \u91cf\u5316\u5931\u6548'],
        ],
        '\u88683-16 LLM\u4e91\u8111\u91cf\u5316\u9002\u7528\u6027\u9a8c\u8bc1'
    )[1]

    # --- Pending items ---
    ref = add_para(doc, ref, '', bold_parts=[
        ('\u5c1a\u7f3a\u7684\u6b63\u5f0f\u8bc4\u4f30\uff08\u5982\u65f6\u95f4\u5141\u8bb8\u8865\u9f50\uff09\uff1a', True),
        ('\u56db\u7b56\u7565\uff08\u542b\u72ec\u7acbDQN\u3001Max-Pressure\uff09\u5728\u76f8\u540c30\u8def\u53e3\u4e09\u573a\u666f\u3001\u540c\u4e00\u79cd\u5b50\u7ec4\u4e0a\u7684\u5b8c\u6574\u5bf9\u6bd4\uff1b\u663e\u8457\u6027\u68c0\u9a8c\u4e0e\u7f6e\u4fe1\u533a\u95f4\uff1b\u5b89\u5168\u6307\u6807\u540c\u53e3\u5f84\u7edf\u8ba1\u3002\u672a\u8865\u9f50\u524d\uff0c\u672c\u6587\u4e0d\u5c06\u5176\u5199\u6210\u201c\u5df2\u5b8c\u6210\u768430\u8def\u53e3\u6700\u7ec8\u7ed3\u8bba\u201d\u3002', False),
    ])[1]

    # Save
    print('Saving...')
    doc.save(DOC)
    print(f'Saved: {DOC}')

if __name__ == '__main__':
    main()
