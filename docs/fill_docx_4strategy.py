#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""回填 docx 3.4 节：四策略全量对比（表+图+正文），并同步 3.1.5/3.4.6 相关文字

数据来源：models/dqn/3scenario_4strategy_totals.json（由 merge_4strategy_report.py 生成）。
图：docs/charts/fig3_14_four_strategies.png（由 chart_4strategy.py 生成，可选）。

修改点：
1. 3.1.5 基线 bullet：Max-Pressure 描述改为四相位实现 + 负面结果指引；
2. 3.4.3 尾部（3.4.4 标题前）插入：导语段 + 表（3场景×4策略）+ 结果解读段 + 图3-14；
3. 3.4.3 安全口径说明段追加 30 路口四策略碰撞合计；
4. 3.4.6 尾部"尚缺正式评估"段改写为"已补齐 + 后续工作"。

幂等：已含"四策略全量对比"导语或目标文字已改写时自动跳过对应步骤。

用法:
    python docs/fill_docx_4strategy.py            # 直接改 docs/系统设计与算法报告 .docx
    python docs/fill_docx_4strategy.py --docx 副本.docx --totals models/dqn/_test_totals.json
"""
from __future__ import annotations

import argparse
import copy
import json
import shutil
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt, Cm

ROOT = Path(__file__).resolve().parent.parent
DEFAULT_DOCX = ROOT / "docs" / "系统设计与算法报告 .docx"
DEFAULT_TOTALS = ROOT / "models" / "dqn" / "3scenario_4strategy_totals.json"
CHART_PNG = ROOT / "docs" / "charts" / "fig3_14_four_strategies.png"

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"

SCEN_LABEL = {"real_peak": "早高峰", "real_offpeak": "平峰", "real_evening": "晚高峰"}
SCEN_ORDER = ["real_peak", "real_offpeak", "real_evening"]
STRAT_LABEL = {"Fixed-Time": "Fixed-Time", "DQN-Original": "共享掩码DQN",
               "Max-Pressure": "Max-Pressure", "Random": "Random"}
STRAT_ORDER = ["Fixed-Time", "DQN-Original", "Max-Pressure", "Random"]

BODY_SIZE = 10.5


def fmt_int(v: float) -> str:
    return f"{v:,.0f}"


def fmt_delta(v: float, base: float) -> str:
    if base == 0:
        return "-"
    return f"{100.0 * (v - base) / base:+.1f}%"


# ---------------- 文本替换 ----------------
def replace_paragraph_text(p, new_text: str, size: float = BODY_SIZE):
    """整体改写段落文本（保留首段样式），字号按正文统一。"""
    # 删除除首个 run 外的所有 run，复用首 run 的字符级格式
    runs = p.runs
    if runs:
        for r in runs[1:]:
            r._r.getparent().remove(r._r)
        runs[0].text = new_text
    else:
        run = p.add_run(new_text)
        run.font.size = Pt(size)


def find_paragraph(doc, marker: str):
    for p in doc.paragraphs:
        if p.text.strip().startswith(marker):
            return p
    return None


# ---------------- 表格 ----------------
def _copy_tbl_borders(source_tbl, target_tbl):
    """把源表格的边框/jc/tblLayout 复制给目标表格（视觉与相邻表一致）。"""
    src_tblPr = source_tbl._tbl.find(W + "tblPr")
    dst_tblPr = target_tbl._tbl.find(W + "tblPr")
    for tag in ("tblBorders", "jc", "tblLayout"):
        el = src_tblPr.find(W + tag)
        if el is not None and dst_tblPr.find(W + tag) is None:
            dst_tblPr.append(copy.deepcopy(el))
    # 宽度自适应
    target_tbl.autofit = True


def _set_cell(cell, text: str, bold: bool = False, size: float = 9.0,
              align: str = "center", name: str = "Calibri"):
    p = cell.paragraphs[0]
    for r in list(p.runs):
        r._r.getparent().remove(r._r)
    run = p.add_run(text)
    run.font.size = Pt(size)
    run.font.name = name
    rPr = run._r.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = rPr.makeelement(qn("w:rFonts"), {})
        rPr.append(rFonts)
    rFonts.set(qn("w:eastAsia"), "宋体")
    run.font.bold = bold
    p.alignment = {"left": 0, "center": 1}.get(align)
    if not text:
        p.add_run("")  # 保证单元格有段落内容


def build_comparison_table(doc, totals: dict):
    """三场景 × 四策略对比表（13 行 × 6 列；Δ% 并入数值列，风格同 30 路口表）"""
    n_rows, n_cols = 13, 6
    tbl = doc.add_table(rows=n_rows, cols=n_cols)

    # 找到文中最接近样式的 30 路口表（正文第 18 张）复制边框
    src = None
    for t in doc.tables:
        if len(t.rows) == 4 and len(t.columns) == 4:
            src = t
            break
    if src is not None:
        _copy_tbl_borders(src, tbl)

    headers = ["场景", "策略", "奖励总和（Δ%）", "等待总和(s)（Δ%）", "碰撞", "胜路口/30"]
    for j, h in enumerate(headers):
        _set_cell(tbl.cell(0, j), h, bold=True)

    # 先做场景列纵向合并，再统一填充（避免合并格残留多段文字）
    for s in range(3):
        merged = tbl.cell(1 + 4 * s, 0).merge(tbl.cell(4 + 4 * s, 0))
        _set_cell(merged, SCEN_LABEL[SCEN_ORDER[s]])

    for s, scen in enumerate(SCEN_ORDER):
        ft = totals[scen]["Fixed-Time"]
        base_reward = ft["reward"]
        base_wait = ft["waiting_time"]
        for k, strat in enumerate(STRAT_ORDER):
            r = 1 + 4 * s + k
            a = totals[scen][strat]
            _set_cell(tbl.cell(r, 1), STRAT_LABEL[strat], align="left")
            reward_cell = (f"{fmt_int(a['reward'])}\n({fmt_delta(a['reward'], base_reward)})"
                           if strat != "Fixed-Time" else fmt_int(a["reward"]))
            wait_cell = (f"{fmt_int(a['waiting_time'])}\n({fmt_delta(a['waiting_time'], base_wait)})"
                         if strat != "Fixed-Time" else fmt_int(a["waiting_time"]))
            _set_cell(tbl.cell(r, 2), reward_cell)
            _set_cell(tbl.cell(r, 3), wait_cell)
            _set_cell(tbl.cell(r, 4), f"{a['collisions']:.0f}")
            _set_cell(tbl.cell(r, 5), "—" if strat == "Fixed-Time" else f"{a['wins']}/30")

    # ---- 固定列宽（cm），避免数字列折行 ----
    _fix_table_widths(tbl, [1.5, 2.5, 2.9, 3.1, 1.3, 1.9])
    return tbl


def _fix_table_widths(tbl, widths_cm):
    """固定布局 + 按列设置 tblGrid / tcW + 行禁止跨页拆分 + 表头行跨页重复。"""
    from docx.oxml import OxmlElement as _OE
    Wns = W
    grid = tbl._tbl.find(Wns + "tblGrid")
    if grid is not None:
        for gc, w in zip(grid.findall(Wns + "gridCol"), widths_cm):
            gc.set(Wns + "w", str(int(w * 567)))
    for i, r in enumerate(tbl.rows):
        for j, w in enumerate(widths_cm):
            tc = r.cells[j]._tc
            tcPr = tc.get_or_add_tcPr()
            tcW = tcPr.find(Wns + "tcW")
            if tcW is None:
                tcW = _OE("w:tcW")
                tcPr.append(tcW)
            tcW.set(Wns + "w", str(int(w * 567)))
            tcW.set(Wns + "type", "dxa")
        # cantSplit：一整行数据不被拆到两页；首行 tblHeader：跨页时重复表头
        trPr = r._tr.find(Wns + "trPr")
        if trPr is None:
            trPr = _OE("w:trPr")
            r._tr.insert(0, trPr)
        if trPr.find(Wns + "cantSplit") is None:
            trPr.append(_OE("w:cantSplit"))
        if i == 0 and trPr.find(Wns + "tblHeader") is None:
            trPr.append(_OE("w:tblHeader"))
    tblPr = tbl._tbl.find(Wns + "tblPr")
    layout = tblPr.find(Wns + "tblLayout")
    if layout is None:
        layout = _OE("w:tblLayout")
        tblPr.append(layout)
    layout.set(Wns + "type", "fixed")
    # 单元格左右边距 0.14 cm，给文字更多可用宽度
    cellMar = tblPr.find(Wns + "tblCellMar")
    if cellMar is None:
        cellMar = _OE("w:tblCellMar")
        tblPr.append(cellMar)
    for side in ("left", "right"):
        el = cellMar.find(Wns + side)
        if el is None:
            el = _OE(f"w:{side}")
            cellMar.append(el)
        el.set(Wns + "w", "80")
        el.set(Wns + "type", "dxa")


def build_analysis_paragraph(totals: dict) -> str:
    """基于 totals 生成一段结果解读（含关键数字与机制说明）。"""
    lines = []
    for scen in SCEN_ORDER:
        ft = totals[scen]["Fixed-Time"]
        dqn = totals[scen]["DQN-Original"]
        mp = totals[scen]["Max-Pressure"]
        rnd = totals[scen]["Random"]
        lines.append(
            f"{SCEN_LABEL[scen]}：共享掩码DQN reward {fmt_delta(dqn['reward'], ft['reward'])}"
            f"（{fmt_int(dqn['reward'])} vs {fmt_int(ft['reward'])}，胜 {dqn['wins']}/30），"
            f"waiting {fmt_delta(dqn['waiting_time'], ft['waiting_time'])}；"
            f"Random（无学习下限）reward {fmt_delta(rnd['reward'], ft['reward'])}；"
            f"Max-Pressure（四相位排队压力）reward {fmt_delta(mp['reward'], ft['reward'])}、"
            f"waiting {fmt_delta(mp['waiting_time'], ft['waiting_time'])}、碰撞 {mp['collisions']:.0f} 起"
        )
    mechanism = ("解读：Max-Pressure 在每个相位切换点贪心选择“放行链路实时停车数之和”最大的相位。"
                 "在“窄路密网”的短车道与独立左转相位结构下，直行方向的停车数通常持续高于左转，"
                 "导致左转相位长期得不到放行（相位饿死），排队进一步向进口道溢流并反噬直行通行，"
                 "最终使该策略在全部三场景中大幅劣于其余策略，碰撞也集中于其切换过程。"
                 "该负面结果说明：仅依赖本路口当前排队信息的贪心相位选择在本类路网中不足以稳定运行，"
                 "这也是本项目采用“需求门控动作掩码 + 学习型共享策略”的动机之一；"
                 "Random 虽无学习能力，但因均匀轮换各相位而不会系统性饿死某类转向，可作为评估下限参考。")
    return "结果解读：" + "；".join(lines) + "。" + mechanism


def insert_block_before(doc, anchor_p, elements):
    """把一组新元素按顺序插入到 anchor_p 之前。elements 为 (kind, obj) 列表。"""
    for kind, obj in elements:
        if kind == "para":
            anchor_p._p.addprevious(obj._p)
        elif kind == "table":
            anchor_p._p.addprevious(obj._tbl)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--docx", type=Path, default=DEFAULT_DOCX)
    ap.add_argument("--totals", type=Path, default=DEFAULT_TOTALS)
    args = ap.parse_args()

    docx_path: Path = args.docx
    totals = json.loads(args.totals.read_text(encoding="utf-8"))

    # 备份（首次运行）
    backup = docx_path.with_suffix(".docx.4strategy.bak")
    if not backup.exists():
        shutil.copy2(docx_path, backup)
        print(f"[BACKUP] {backup}")

    doc = Document(str(docx_path))
    changed = []

    # ---- 1) 3.1.5 Max-Pressure bullet 改写 ----
    p202 = find_paragraph(doc, "• Max-Pressure感应控制")
    if p202 is not None:
        replace_paragraph_text(
            p202,
            "• Max-Pressure（四相位排队压力）感应控制：相位压力 = 该相位放行链路上的实时停车车辆数之和"
            "（覆盖直行/左转全部四相位），每个切换点选择压力最大的相位，遵守最小绿灯15s与黄灯过渡；"
            "在本路网30路口全量评估中呈系统性负面结果（左转相位饿死→排队溢流，见3.4.3），"
            "故仅作为无学习参照基线，不推荐为最终控制方案。",
        )
        changed.append("3.1.5 Max-Pressure 描述")

    # ---- 2) 3.4.3 安全口径说明段追加 ----
    p265 = find_paragraph(doc, "安全口径说明")
    if p265 is not None and "下碰撞合计" not in p265.text:
        stats = []
        for strat in STRAT_ORDER:
            vals = [f"{totals[sc][strat]['collisions']:.0f}" for sc in SCEN_ORDER]
            stats.append(f"{STRAT_LABEL[strat]} {vals[0]}/{vals[1]}/{vals[2]}")
        extra = ("30路口全量口径（3回合×720步）下碰撞合计（早高峰/平峰/晚高峰）："
                 + "；".join(stats) + "（Fixed-Time 全程为0，碰撞集中在参照基线上，详见3.4.3对比表）。")
        run = p265.add_run(extra)
        run.font.size = Pt(BODY_SIZE)
        changed.append("3.4.3 安全口径说明追加")

    # ---- 3) 3.4.4 标题前插入四策略区块 ----
    anchor = find_paragraph(doc, "3.4.4 轻量化与部署性能验证")
    if anchor is not None and find_paragraph(doc, "四策略全量对比") is None:
        # 导语段（粗体前缀）
        lead = doc.add_paragraph()
        lead.style = doc.styles["Normal"]
        r1 = lead.add_run("四策略全量对比（3回合×720步，同种子42–44）：")
        r1.font.size = Pt(BODY_SIZE)
        r1.font.bold = True
        r2 = lead.add_run("在相同环境协议下补齐 Max-Pressure（四相位排队压力，见3.1.5）与 "
                          "Random 的30路口三场景评估，reward/waiting 聚合口径与上文一致（Σ路口 mean），"
                          "策略行内给出对 Fixed-Time 的变化率与胜路口统计。")
        r2.font.size = Pt(BODY_SIZE)

        tbl = build_comparison_table(doc, totals)

        analysis = doc.add_paragraph()
        analysis.style = doc.styles["Normal"]
        a_run = analysis.add_run(build_analysis_paragraph(totals))
        a_run.font.size = Pt(BODY_SIZE)

        elements = [("para", lead), ("table", tbl), ("para", analysis)]
        extra_note = ""

        if CHART_PNG.exists():
            img_para = doc.add_paragraph()
            img_para.alignment = 1  # center
            run = img_para.add_run()
            run.add_picture(str(CHART_PNG), width=Cm(15.2))
            cap = doc.add_paragraph()
            cap.style = doc.styles["Normal"]
            cap.alignment = 1
            cr = cap.add_run("图3-14 四策略×30路口全量对比（reward 为 Σ路口 mean；waiting 数量级差异大，采用对数轴）")
            cr.font.size = Pt(9)
            elements.append(("para", img_para))
            elements.append(("para", cap))
            extra_note = " + 图3-14"

        insert_block_before(doc, anchor, elements)
        changed.append("3.4.3 插入四策略对比表 + 解读" + extra_note)

    # ---- 4) 3.4.6 尾部"尚缺正式评估"改写 ----
    p278 = find_paragraph(doc, "尚缺的正式评估")
    if p278 is not None:
        replace_paragraph_text(
            p278,
            "正式评估补充说明：四策略（Fixed-Time / 共享掩码DQN / Max-Pressure / Random）已按"
            "相同30路口三场景、同一种子组（42–44）完成同口径完整对比（见3.4.3对比表）。"
            "其中 Max-Pressure 因左转相位饿死呈系统性负面结果，Random 作为无学习下限参考；"
            "碰撞等安全指标已同口径统计。本文据此给出30路口最终结论；"
            "显著性检验与置信区间、更多随机种子的补充留作后续工作。",
        )
        changed.append("3.4.6 尚缺评估文字改写")

    doc.save(str(docx_path))
    print(f"[SAVE] {docx_path}")
    print("已执行修改：" + ("；".join(changed) if changed else "（无——文档已是目标状态）"))


if __name__ == "__main__":
    main()
